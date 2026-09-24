"""Pruebas de los resúmenes con IA usando un cliente de Claude simulado (sin coste)."""

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from uji_sync.ai import (
    SUMMARIES_DIR, AIError, AIFatalError, AIProcessor, Summarizer, Unsupported,
    document_block, needs_api, pending, summary_relpath,
)
from uji_sync.registry import FileRecord, Registry

DATA = {
    "titulo": "Introducción al cálculo",
    "tipo_documento": "apuntes",
    "en_una_frase": "Límites y derivadas básicas.",
    "resumen": "El tema presenta los **límites** y las derivadas.",
    "puntos_clave": ["Definición de límite", "Regla de la cadena"],
    "conceptos": [{"termino": "Derivada", "definicion": "Tasa de cambio instantánea."}],
    "preguntas_repaso": ["¿Qué es un límite?"],
}


class FakeMessages:
    def __init__(self, stop_reason="end_turn", data=DATA):
        self.calls = []
        self.stop_reason = stop_reason
        self.data = data

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(
            stop_reason=self.stop_reason,
            model=kwargs["model"],
            content=[SimpleNamespace(type="thinking", thinking=""),
                     SimpleNamespace(type="text", text=json.dumps(self.data))],
            usage=SimpleNamespace(input_tokens=10_000, output_tokens=2_000),
        )


def fake_client(**kw):
    msgs = FakeMessages(**kw)
    return SimpleNamespace(beta=SimpleNamespace(messages=msgs)), msgs


# ------------------------------------------------------------------ documentos
def test_document_blocks(tmp_path):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4 hola")
    b = document_block(pdf, "a.pdf")
    assert b["source"]["type"] == "base64" and b["source"]["media_type"] == "application/pdf"

    import docx
    d = docx.Document()
    d.add_paragraph("Tema 1: límites")
    d.save(tmp_path / "b.docx")
    b = document_block(tmp_path / "b.docx", "b.docx")
    assert b["source"] == {"type": "text", "media_type": "text/plain", "data": "Tema 1: límites"}

    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Derivadas"
    slide.notes_slide.notes_text_frame.text = "Explicar con ejemplos"
    prs.save(tmp_path / "c.pptx")
    text = document_block(tmp_path / "c.pptx", "c.pptx")["source"]["data"]
    assert "Diapositiva 1" in text and "Derivadas" in text and "Notas: Explicar con ejemplos" in text

    (tmp_path / "vacio.txt").write_text("  ")
    with pytest.raises(Unsupported):
        document_block(tmp_path / "vacio.txt", "vacio.txt")
    (tmp_path / "x.xlsx").write_bytes(b"x")
    with pytest.raises(Unsupported):
        document_block(tmp_path / "x.xlsx", "x.xlsx")


# ------------------------------------------------------------------ Summarizer
def test_summarizer_request(tmp_path):
    pdf = tmp_path / "Tema 1.pdf"
    pdf.write_bytes(b"%PDF-1.4 hola")
    client, msgs = fake_client()
    s = Summarizer(client=client).summarize(pdf, "Matemáticas", "Tema 1")
    assert s.data == DATA and (s.input_tokens, s.output_tokens) == (10_000, 2_000)
    req = msgs.calls[0]
    assert req["model"] == "claude-opus-5"
    assert req["fallbacks"] == "default" and req["betas"] == ["server-side-fallback-2026-07-01"]
    assert req["thinking"] == {"type": "adaptive"}
    assert req["output_config"]["format"]["type"] == "json_schema"
    content = req["messages"][0]["content"]
    assert content[0]["type"] == "document" and "Matemáticas" in content[1]["text"]

    client, msgs = fake_client()
    Summarizer(client=client, model="claude-sonnet-5").summarize(pdf, "M", "T")
    assert "fallbacks" not in msgs.calls[0]


@pytest.mark.parametrize("reason", ["refusal", "max_tokens"])
def test_summarizer_bad_stop_reason(tmp_path, reason):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF")
    client, _ = fake_client(stop_reason=reason)
    with pytest.raises(AIError):
        Summarizer(client=client).summarize(pdf, "M", "T")


# ------------------------------------------------------------------ proceso
class FakeSummarizer:
    def __init__(self, fail_on=None, fatal_on=None):
        self.calls = []
        self.fail_on, self.fatal_on = fail_on, fatal_on

    def summarize(self, path, course, section):
        self.calls.append(path.name)
        if path.name == self.fatal_on:
            raise AIFatalError("clave no válida")
        if path.name == self.fail_on:
            raise AIError("rechazado")
        from uji_sync.ai import Summary
        return Summary({**DATA, "en_una_frase": f"Idea de {path.name}"}, "claude-opus-5", 10_000, 2_000)


def add_file(reg, root, local, content, course_id=1, course="Matemáticas", section="Tema 1"):
    import hashlib
    path = root / local
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    rec = FileRecord(key=f"k:{local}", kind="file", course_id=course_id, course_name=course,
                     section=section, name=Path(local).name, url="u", local_path=local,
                     sha256=hashlib.sha256(content).hexdigest(), size=len(content),
                     first_seen="t", last_seen="t", last_changed="t")
    reg.upsert(rec)
    return rec


def test_processor_writes_summaries_and_news(tmp_path):
    root = tmp_path / "UJI"
    reg = Registry(tmp_path / "r.db")
    a = add_file(reg, root, "Matemáticas/01 - Tema 1/Tema 1.pdf", b"%PDF a")
    b = add_file(reg, root, "Matemáticas/01 - Tema 1/Ejercicios.docx", b"docx b")
    x = add_file(reg, root, "Matemáticas/01 - Tema 1/datos.xlsx", b"xlsx")
    fake = FakeSummarizer()
    proc = AIProcessor(root, reg, summarizer=fake, log=lambda _: None)

    assert [r.key for r in pending(reg, root)] == [b.key, a.key]  # el .xlsx no se resume
    res = proc.process([a, b, x])
    assert fake.calls == ["Tema 1.pdf", "Ejercicios.docx"]
    assert res.courses[1].summarized == 2
    assert res.cost == pytest.approx(2 * (10_000 * 5 + 2_000 * 25) / 1e6)
    assert "Matemáticas → 2 resumidos" in res.text()

    md = root / summary_relpath(a.local_path)
    assert md == root / "Matemáticas" / SUMMARIES_DIR / "01 - Tema 1" / "Tema 1.pdf - resumen.md"
    text = md.read_text(encoding="utf-8")
    assert "# Introducción al cálculo" in text
    assert "[Tema 1.pdf](<../../01 - Tema 1/Tema 1.pdf>)" in text  # enlace al original
    assert "- **Derivada**: Tasa de cambio instantánea." in text
    news = (root / "Matemáticas" / SUMMARIES_DIR / "Novedades.md").read_text(encoding="utf-8")
    assert news.startswith("# Novedades · Matemáticas")
    assert "Idea de Tema 1.pdf" in news and "<01 - Tema 1/Tema 1.pdf - resumen.md>" in news
    assert pending(reg, root) == []

    # Nada se paga dos veces; si se borra un .md se regenera sin llamar a la API
    md.unlink()
    assert [r.key for r in pending(reg, root)] == [a.key] and not needs_api(reg, a)
    res = proc.process([a])
    assert fake.calls == ["Tema 1.pdf", "Ejercicios.docx"] and res.cost == 0
    assert res.courses[1].reused == 1 and md.exists()
    news = (root / "Matemáticas" / SUMMARIES_DIR / "Novedades.md").read_text(encoding="utf-8")
    assert news.count("## ") == 2  # la ejecución nueva se añade arriba


def test_processor_errors(tmp_path):
    root = tmp_path / "UJI"
    reg = Registry(tmp_path / "r.db")
    a = add_file(reg, root, "M/01/a.pdf", b"a")
    b = add_file(reg, root, "M/01/b.pdf", b"b")
    c = add_file(reg, root, "M/01/c.pdf", b"c")
    proc = AIProcessor(root, reg, summarizer=FakeSummarizer(fail_on="a.pdf", fatal_on="b.pdf"),
                       log=lambda _: None)
    res = proc.process([a, b, c])
    assert res.courses[1].errors == ["a.pdf: rechazado"]
    assert res.fatal == "clave no válida" and res.courses[1].summarized == 0  # c no se procesa
    assert "Se detuvo" in res.text()
