"""Resúmenes con IA (Claude) de los materiales descargados.

Para cada archivo nuevo o cambiado:
  1. Se prepara el documento: los PDF se envían tal cual (Claude ve texto, tablas
     e imágenes); de Word y PowerPoint se extrae el texto.
  2. Claude devuelve un JSON con un esquema fijo (salidas estructuradas).
  3. Se guarda en el registro, indexado por el SHA-256 del contenido, y se
     escribe un .md en <Asignatura>/_resumenes_IA/ con la misma estructura de temas.

Un mismo contenido nunca se envía dos veces: si el resumen ya existe en el
registro, solo se vuelve a escribir el .md (sin coste).
"""

from __future__ import annotations

import base64
import json
import posixpath
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Callable

import anthropic

from .apikey import get_api_key
from .registry import FileRecord, Registry

DEFAULT_MODEL = "claude-opus-5"
MODELS = {
    "claude-opus-5": "Claude Opus 5 (mejor calidad)",
    "claude-sonnet-5": "Claude Sonnet 5 (más barato)",
}
# Precios en US$ por millón de tokens (entrada, salida), para estimar el coste.
PRICES = {
    "claude-opus-5": (5.0, 25.0),
    "claude-opus-4-8": (5.0, 25.0),  # modelo de reserva si Opus 5 rechaza un documento
    "claude-sonnet-5": (2.0, 10.0),
}
SUMMARIES_DIR = "_resumenes_IA"
NEWS_FILE = "Novedades.md"
SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}
MAX_PDF_BYTES = 22 * 1024 * 1024   # el límite de la API es 32 MB por petición (en base64)
MAX_TEXT_CHARS = 1_500_000         # ~400K tokens; más largo se omite en vez de recortarlo

SYSTEM_PROMPT = """\
Eres un asistente de estudio para un estudiante de la Universitat Jaume I.
Recibirás un material de una de sus asignaturas (apuntes, diapositivas,
ejercicios, guías...). Escribe en español, aunque el documento esté en
valenciano, inglés u otro idioma, y respeta la terminología técnica.

Sé fiel al documento: no añadas contenido que no aparezca en él. Si el
documento es un enunciado de ejercicios o un examen, resume qué se pide y qué
conocimientos hacen falta, sin resolverlo. Si apenas tiene contenido (una
portada, un horario), dilo brevemente en el resumen y deja vacías las listas
que no apliquen."""

SUMMARY_SCHEMA = {
    "type": "object",
    "properties": {
        "titulo": {"type": "string", "description": "Título descriptivo del documento"},
        "tipo_documento": {
            "type": "string",
            "enum": ["apuntes", "diapositivas", "ejercicios", "práctica", "examen",
                     "guía docente", "información", "otro"],
        },
        "en_una_frase": {"type": "string", "description": "La idea principal en una frase"},
        "resumen": {"type": "string", "description": "Resumen en Markdown, 1 a 5 párrafos"},
        "puntos_clave": {"type": "array", "items": {"type": "string"}},
        "conceptos": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"termino": {"type": "string"}, "definicion": {"type": "string"}},
                "required": ["termino", "definicion"],
                "additionalProperties": False,
            },
        },
        "preguntas_repaso": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["titulo", "tipo_documento", "en_una_frase", "resumen", "puntos_clave",
                 "conceptos", "preguntas_repaso"],
    "additionalProperties": False,
}


class AIError(Exception):
    """Error con un documento concreto: se anota y se sigue con el siguiente."""


class AIFatalError(Exception):
    """Error que impide seguir (clave no válida, sin conexión, límite de uso)."""


class Unsupported(Exception):
    pass


# ---------------------------------------------------------------- documentos
def docx_text(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def pptx_text(path: Path) -> str:
    from pptx import Presentation

    parts = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts.append(f"--- Diapositiva {n} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notas: {notes}")
    return "\n".join(parts)


def document_block(path: Path, title: str) -> dict:
    ext = path.suffix.lower()
    if ext == ".pdf":
        data = path.read_bytes()
        if len(data) > MAX_PDF_BYTES:
            raise Unsupported("PDF demasiado grande para la IA")
        return {
            "type": "document", "title": title,
            "source": {"type": "base64", "media_type": "application/pdf",
                       "data": base64.standard_b64encode(data).decode("ascii")},
        }
    if ext == ".docx":
        text = docx_text(path)
    elif ext == ".pptx":
        text = pptx_text(path)
    elif ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise Unsupported(f"formato {ext} no compatible")
    if not text.strip():
        raise Unsupported("no contiene texto")
    if len(text) > MAX_TEXT_CHARS:
        raise Unsupported("documento demasiado largo para la IA")
    return {
        "type": "document", "title": title,
        "source": {"type": "text", "media_type": "text/plain", "data": text},
    }


# ---------------------------------------------------------------- Claude
@dataclass
class Summary:
    data: dict
    model: str
    input_tokens: int
    output_tokens: int


def make_client() -> anthropic.Anthropic:
    key = get_api_key()
    # Sin clave guardada, el SDK usa ANTHROPIC_API_KEY u otras credenciales del entorno.
    return anthropic.Anthropic(api_key=key) if key else anthropic.Anthropic()


class Summarizer:
    def __init__(self, client=None, model: str = DEFAULT_MODEL, effort: str = "medium"):
        self.client = client or make_client()
        self.model = model
        self.effort = effort

    def summarize(self, path: Path, course: str, section: str) -> Summary:
        block = document_block(path, path.name)
        request = dict(
            model=self.model,
            max_tokens=16000,
            system=SYSTEM_PROMPT,
            thinking={"type": "adaptive"},
            output_config={
                "effort": self.effort,
                "format": {"type": "json_schema", "schema": SUMMARY_SCHEMA},
            },
            messages=[{"role": "user", "content": [
                block,
                {"type": "text", "text": f"Asignatura: {course}\nSección: {section}\n"
                                         f"Archivo: {path.name}\n\nResume este material."},
            ]}],
        )
        if self.model == "claude-opus-5":
            # Si el filtro de seguridad rechaza un documento, se reintenta en el
            # servidor con el modelo de reserva recomendado.
            request.update(betas=["server-side-fallback-2026-07-01"], fallbacks="default")
        try:
            resp = self.client.beta.messages.create(**request)
        except anthropic.BadRequestError as e:
            raise AIError(f"la API rechazó el documento: {e.message}") from e
        except anthropic.AuthenticationError as e:
            raise AIFatalError("La clave de API de Claude no es válida.") from e
        except anthropic.PermissionDeniedError as e:
            raise AIFatalError("La clave de API no tiene permiso para usar este modelo.") from e
        except anthropic.RateLimitError as e:
            raise AIFatalError("Se ha alcanzado el límite de uso de la API. Prueba más tarde.") from e
        except anthropic.APIStatusError as e:
            if e.status_code >= 500:
                raise AIError(f"error temporal del servicio ({e.status_code})") from e
            raise AIFatalError(f"Error de la API ({e.status_code}): {e.message}") from e
        except anthropic.APIConnectionError as e:
            raise AIFatalError("No hay conexión con la API de Claude.") from e
        except anthropic.AnthropicError as e:
            raise AIFatalError(f"No se pudo usar la IA: {e}") from e

        if resp.stop_reason == "refusal":
            raise AIError("la IA no ha podido procesar este documento")
        if resp.stop_reason == "max_tokens":
            raise AIError("el resumen ha quedado incompleto")
        text = next((b.text for b in resp.content if b.type == "text"), "")
        try:
            data = json.loads(text)
        except ValueError as e:
            raise AIError("respuesta de la IA no válida") from e
        return Summary(data, resp.model or self.model,
                       resp.usage.input_tokens, resp.usage.output_tokens)


def estimate_cost(model: str, input_tokens: int, output_tokens: int) -> float:
    p_in, p_out = PRICES.get(model, PRICES[DEFAULT_MODEL])
    return input_tokens * p_in / 1e6 + output_tokens * p_out / 1e6


# ---------------------------------------------------------------- salida .md
def summary_relpath(local_path: str) -> str:
    """Matemáticas/01 - Tema 1/Tema 1.pdf -> Matemáticas/_resumenes_IA/01 - Tema 1/Tema 1.pdf - resumen.md"""
    p = PurePosixPath(local_path)
    return str(PurePosixPath(p.parts[0], SUMMARIES_DIR, *p.parts[1:-1], f"{p.name} - resumen.md"))


def _link(from_file: str, to_file: str) -> str:
    rel = posixpath.relpath(to_file, posixpath.dirname(from_file))
    return f"<{rel}>"


def render_markdown(data: dict, rec: FileRecord, summary_path: str, model: str) -> str:
    lines = [
        f"# {data['titulo']}",
        "",
        f"*{rec.course_name} · {rec.section or 'General'} · "
        f"[{PurePosixPath(rec.local_path).name}]({_link(summary_path, rec.local_path)}) · "
        f"{data['tipo_documento']}*",
        "",
        f"> **En una frase:** {data['en_una_frase']}",
        "",
        "## Resumen",
        "",
        data["resumen"].strip(),
    ]
    if data["puntos_clave"]:
        lines += ["", "## Puntos clave", ""] + [f"- {p}" for p in data["puntos_clave"]]
    if data["conceptos"]:
        lines += ["", "## Conceptos", ""] + [
            f"- **{c['termino']}**: {c['definicion']}" for c in data["conceptos"]]
    if data["preguntas_repaso"]:
        lines += ["", "## Preguntas de repaso", ""] + [
            f"{i}. {q}" for i, q in enumerate(data["preguntas_repaso"], 1)]
    lines += ["", "---", "",
              f"_Resumen generado por IA ({model}). Puede contener errores: "
              "consulta siempre el material original._", ""]
    return "\n".join(lines)


# ---------------------------------------------------------------- proceso
@dataclass
class CourseAIResult:
    course_name: str
    summarized: int = 0     # resúmenes nuevos (con coste)
    reused: int = 0         # ya resumidos antes (sin coste)
    skipped: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


@dataclass
class AIRunResult:
    courses: dict[int, CourseAIResult] = field(default_factory=dict)
    cost: float = 0.0
    fatal: str | None = None

    def text(self) -> str:
        if not self.courses and not self.fatal:
            return ""
        lines = ["Resúmenes con IA", ""]
        for r in self.courses.values():
            parts = [f"{r.summarized} {'resumido' if r.summarized == 1 else 'resumidos'}"]
            if r.reused:
                parts.append(f"{r.reused} ya resumidos")
            if r.errors:
                parts.append(f"{len(r.errors)} errores")
            lines.append(f"{r.course_name} → " + ", ".join(parts))
        lines += ["", f"Coste aproximado: {self.cost:.2f} US$".replace(".", ",")]
        if self.fatal:
            lines += ["", f"⚠ Se detuvo: {self.fatal}"]
        return "\n".join(lines)


def is_supported(rec: FileRecord) -> bool:
    return PurePosixPath(rec.local_path).suffix.lower() in SUPPORTED


def needs_api(registry: Registry, rec: FileRecord) -> bool:
    return is_supported(rec) and registry.get_summary(rec.sha256) is None


def pending(registry: Registry, root: Path, course_ids: list[int] | None = None) -> list[FileRecord]:
    """Archivos sin resumen (o cuyo .md se ha borrado)."""
    return [
        r for r in registry.files(course_ids)
        if is_supported(r) and (registry.get_summary(r.sha256) is None
                                or not (root / summary_relpath(r.local_path)).exists())
    ]


class AIProcessor:
    def __init__(
        self,
        root: Path,
        registry: Registry,
        summarizer: Summarizer | None = None,
        model: str = DEFAULT_MODEL,
        log: Callable[[str], None] = print,
        should_stop: Callable[[], bool] = lambda: False,
    ):
        self.root = root
        self.registry = registry
        self._summarizer = summarizer
        self.model = model
        self.log = log
        self.should_stop = should_stop

    @property
    def summarizer(self) -> Summarizer:
        if self._summarizer is None:  # solo se crea el cliente si hace falta llamar a la API
            try:
                self._summarizer = Summarizer(model=self.model)
            except anthropic.AnthropicError as e:
                raise AIFatalError(f"No se pudo iniciar la IA: {e}") from e
        return self._summarizer

    def process(self, records: list[FileRecord]) -> AIRunResult:
        result = AIRunResult()
        news: dict[int, list[tuple[FileRecord, dict, str]]] = {}
        for rec in records:
            if self.should_stop() or result.fatal:
                break
            if not is_supported(rec):
                continue
            cres = result.courses.setdefault(rec.course_id, CourseAIResult(rec.course_name))
            name = PurePosixPath(rec.local_path).name
            data = self.registry.get_summary(rec.sha256)
            if data is not None:
                cres.reused += 1
                model = data.get("_model", "")
            else:
                self.log(f"  🤖 Resumiendo {rec.local_path}…")
                try:
                    s = self.summarizer.summarize(self.root / rec.local_path,
                                                  rec.course_name, rec.section)
                except Unsupported as e:
                    cres.skipped.append(name)
                    self.log(f"  – {name}: {e}")
                    continue
                except AIError as e:
                    cres.errors.append(f"{name}: {e}")
                    self.log(f"  ✗ {name}: {e}")
                    continue
                except AIFatalError as e:
                    result.fatal = str(e)
                    self.log(f"  ✗ {e}")
                    break
                except OSError as e:
                    cres.errors.append(f"{name}: {e}")
                    self.log(f"  ✗ {name}: no se pudo leer ({e})")
                    continue
                self.registry.save_summary(rec.sha256, s.model, s.data,
                                           s.input_tokens, s.output_tokens)
                result.cost += estimate_cost(s.model, s.input_tokens, s.output_tokens)
                cres.summarized += 1
                data, model = s.data, s.model
            spath = summary_relpath(rec.local_path)
            self._write(spath, render_markdown(data, rec, spath, model))
            news.setdefault(rec.course_id, []).append((rec, data, spath))
        for items in news.values():
            self._update_news(items)
        return result

    def _write(self, rel: str, text: str) -> None:
        path = self.root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")

    def _update_news(self, items: list[tuple[FileRecord, dict, str]]) -> None:
        """Añade arriba de <Asignatura>/_resumenes_IA/Novedades.md lo de esta ejecución."""
        course = items[0][0].course_name
        course_dir = PurePosixPath(items[0][0].local_path).parts[0]
        rel = str(PurePosixPath(course_dir, SUMMARIES_DIR, NEWS_FILE))
        header = f"# Novedades · {course}\n"
        section = [f"\n## {datetime.now().strftime('%d/%m/%Y %H:%M')}\n"]
        for rec, data, spath in items:
            section.append(
                f"- **[{PurePosixPath(rec.local_path).name}]({_link(rel, spath)})** "
                f"({rec.section or 'General'}): {data['en_una_frase']}"
            )
        path = self.root / rel
        old = path.read_text(encoding="utf-8") if path.exists() else header
        body = old[len(header):] if old.startswith(header) else old
        self._write(rel, header + "\n".join(section) + "\n" + body)
