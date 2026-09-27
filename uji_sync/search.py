"""Búsqueda dentro de tus documentos (Fase 4).

- Extrae el texto de PDF, Word, PowerPoint y textos, recordando la página o
  diapositiva. Gratis y en tu PC.
- Índice por palabras (SQLite FTS5): ignora tildes y mayúsculas.
- Búsqueda por significado con un modelo multilingüe que se ejecuta en tu PC
  (fastembed). Si no está disponible, se usa solo la búsqueda por palabras.
- Fotos y PDF escaneados no tienen texto: con la IA activada, Claude los lee
  una vez (OCR) y su texto pasa al índice.

El índice vive junto al registro, en %LOCALAPPDATA%\\UJISync\\registros\\.
"""

from __future__ import annotations

import re
import sqlite3
import threading
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Callable

from .ai import AIFatalError
from .config import app_dir, registry_path
from .library import course_materials, list_courses

TEXT_EXTS = {".pdf", ".docx", ".pptx", ".txt", ".md"}
OCR_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".pdf"}
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150
MIN_PDF_CHARS_PER_PAGE = 25   # por debajo, el PDF se considera escaneado
EMBED_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
MIN_SIMILARITY = 0.35  # por debajo, el fragmento se considera poco relacionado
HL_OPEN, HL_CLOSE = "\u0002", "\u0003"  # marcas de resaltado (la web las convierte)
STOPWORDS = set("""a al algo ante bajo como con cual cuales cuando de del donde el ella en entre
es esa ese eso esta este esto ha hay la las le lo los mas me mi mis muy no o para pero por que
qué se sin sobre su sus te tu un una uno unos unas y ya yo""".split())

SCHEMA = """
CREATE TABLE IF NOT EXISTS docs (
    path      TEXT PRIMARY KEY,
    signature TEXT NOT NULL,     -- tamaño + fecha de modificación
    status    TEXT NOT NULL,     -- ok | sin_texto | error
    source    TEXT NOT NULL,     -- texto | ocr
    detail    TEXT NOT NULL,
    indexed   TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS chunks (
    id   INTEGER PRIMARY KEY,
    path TEXT NOT NULL,
    loc  TEXT NOT NULL,          -- «pág. 3», «diapositiva 5» o ""
    text TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS chunks_path ON chunks(path);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
    text, content='chunks', content_rowid='id', tokenize='unicode61 remove_diacritics 2'
);
CREATE TABLE IF NOT EXISTS vectors (
    chunk_id INTEGER PRIMARY KEY,
    vec      BLOB NOT NULL
);
CREATE TABLE IF NOT EXISTS ocr (
    signature TEXT PRIMARY KEY,  -- ruta + tamaño + fecha: el mismo archivo no se paga dos veces
    text      TEXT NOT NULL,
    created   TEXT NOT NULL
);
"""


def search_db_path(root: Path) -> Path:
    reg = registry_path(root)
    return reg.with_name(reg.stem + "-busqueda.db")


# ---------------------------------------------------------------- extracción
class NeedsOCR(Exception):
    """El documento no tiene texto (foto o escaneado)."""


def extract_segments(path: Path) -> list[tuple[str, str]]:
    """Lista de (ubicación, texto). Lanza NeedsOCR si no hay texto que extraer."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [(f"pág. {i}", (p.extract_text() or "").strip()) for i, p in enumerate(reader.pages, 1)]
        total = sum(len(t) for _, t in pages)
        if total < MIN_PDF_CHARS_PER_PAGE * max(1, len(pages)):
            raise NeedsOCR("PDF escaneado (sin texto)")
        return [(loc, t) for loc, t in pages if t]
    if ext == ".docx":
        from .ai import docx_text

        return [("", docx_text(path))]
    if ext == ".pptx":
        from pptx import Presentation

        out = []
        for n, slide in enumerate(Presentation(str(path)).slides, 1):
            parts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
            for s in slide.shapes:
                if getattr(s, "has_table", False) and s.has_table:
                    parts += [" | ".join(c.text for c in row.cells) for row in s.table.rows]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("Notas: " + slide.notes_slide.notes_text_frame.text)
            if parts:
                out.append((f"diapositiva {n}", "\n".join(parts)))
        return out
    if ext in (".txt", ".md"):
        return [("", path.read_text(encoding="utf-8", errors="replace"))]
    if ext in OCR_EXTS:
        raise NeedsOCR("imagen")
    return []


def chunk(segments: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """Trocea el texto en fragmentos solapados, sin perder la ubicación."""
    out = []
    for loc, text in segments:
        text = re.sub(r"[ \t]+", " ", text).strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + CHUNK_CHARS)
            if end < len(text):  # cortar en un final de frase o de línea si es posible
                cut = max(text.rfind(". ", start, end), text.rfind("\n", start, end))
                if cut > start + CHUNK_CHARS // 2:
                    end = cut + 1
            piece = text[start:end].strip()
            if piece:
                out.append((loc, piece))
            if end >= len(text):
                break
            start = max(end - CHUNK_OVERLAP, start + 1)
    return out


# ---------------------------------------------------------------- significado
class LocalEmbedder:
    """Modelo multilingüe que se ejecuta en tu PC (se descarga la primera vez)."""

    def __init__(self, model: str = EMBED_MODEL):
        self.model_name = model
        self._model = None
        self.error: str | None = None
        self._lock = threading.Lock()

    def available(self) -> bool:
        try:
            import fastembed  # noqa: F401
        except ImportError:
            self.error = "falta el paquete fastembed (pip install -r requirements.txt)"
            return False
        return self.error is None

    def embed(self, texts: list[str]):
        import numpy as np

        with self._lock:
            if self._model is None:
                try:
                    from fastembed import TextEmbedding

                    self._model = TextEmbedding(self.model_name, cache_dir=str(app_dir() / "modelos"))
                except Exception as e:  # sin internet la primera vez, etc.
                    self.error = f"no se pudo cargar el modelo: {e}"
                    raise
            vecs = np.array(list(self._model.embed(texts)), dtype="float32")
        norms = np.linalg.norm(vecs, axis=1, keepdims=True)
        return vecs / np.maximum(norms, 1e-9)


# ---------------------------------------------------------------- índice
@dataclass
class IndexStats:
    indexed: int = 0
    removed: int = 0
    no_text: list[str] = field(default_factory=list)
    ocr_done: int = 0
    errors: list[str] = field(default_factory=list)
    cost: float = 0.0
    semantic_error: str | None = None

    def text(self) -> str:
        lines = ["Índice de búsqueda actualizado", "",
                 f"{self.indexed} documentos indexados, {self.removed} quitados"]
        if self.ocr_done:
            cost = f"{self.cost:.2f}".replace(".", ",")
            lines.append(f"{self.ocr_done} fotos o escaneados leídos con IA (coste aprox. {cost} US$)")
        if self.no_text:
            lines.append(f"{len(self.no_text)} fotos o escaneados sin leer (se buscan por su título)")
        if self.errors:
            lines.append(f"{len(self.errors)} con errores")
        if self.semantic_error:
            lines.append(f"Búsqueda por significado no disponible: {self.semantic_error}")
        return "\n".join(lines)


def _signature(path: Path) -> str:
    st = path.stat()
    return f"{st.st_size}:{st.st_mtime_ns}"


class SearchIndex:
    def __init__(self, root: Path, embedder: LocalEmbedder | None = None, db_path: Path | None = None):
        self.root = root
        self.embedder = embedder if embedder is not None else LocalEmbedder()
        path = db_path or search_db_path(root)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(path)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)
        self._matrix = None  # vectores en memoria para buscar por significado

    def close(self) -> None:
        self.conn.close()

    # ---------------------------------------------------------- estado
    def _materials(self) -> dict[str, Path]:
        return {rel: self.root / rel for c in list_courses(self.root) for rel in course_materials(self.root, c)
                if PurePosixPath(rel).suffix.lower() in TEXT_EXTS | OCR_EXTS}

    def status(self) -> dict:
        docs = {r["path"]: dict(r) for r in self.conn.execute("SELECT * FROM docs")}
        materials = self._materials()
        pending = [rel for rel, p in materials.items()
                   if rel not in docs or docs[rel]["signature"] != _signature(p)]
        no_text = [rel for rel, d in docs.items() if d["status"] == "sin_texto" and rel in materials]
        chunks = self.conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        vectors = self.conn.execute("SELECT COUNT(*) FROM vectors").fetchone()[0]
        return {"documents": len(materials), "indexed": len(materials) - len(pending),
                "pending": len(pending), "no_text": len(no_text), "chunks": chunks,
                "semantic": {"available": self.embedder.available(), "vectors": vectors,
                             "error": self.embedder.error}}

    def ocr_candidates(self) -> list[str]:
        """Fotos y escaneados indexados sin texto (se pueden leer con IA)."""
        materials = self._materials()
        return [r["path"] for r in self.conn.execute("SELECT * FROM docs WHERE status = 'sin_texto'")
                if r["path"] in materials and r["signature"] == _signature(materials[r["path"]])]

    def _cached_ocr(self, rel: str, path: Path) -> str | None:
        row = self.conn.execute("SELECT text FROM ocr WHERE signature = ?",
                                (f"{rel}|{_signature(path)}",)).fetchone()
        return row["text"] if row else None

    # ---------------------------------------------------------- actualizar
    def update(self, log: Callable[[str], None] = lambda _: None,
               ocr: Callable[[Path], tuple[str, float]] | None = None,
               should_stop: Callable[[], bool] = lambda: False) -> IndexStats:
        """Indexa lo nuevo o cambiado. `ocr(path) -> (texto, coste)` lee fotos con IA."""
        stats = IndexStats()
        docs = {r["path"]: dict(r) for r in self.conn.execute("SELECT * FROM docs")}
        materials = self._materials()
        for rel in [r for r in docs if r not in materials]:
            self._delete(rel)
            stats.removed += 1
        for rel, path in materials.items():
            if should_stop():
                break
            sig = _signature(path)
            d = docs.get(rel)
            retry_ocr = d is not None and d["status"] == "sin_texto" and ocr is not None
            if d is not None and d["signature"] == sig and not retry_ocr:
                continue
            source, status, detail = "texto", "ok", ""
            try:
                try:
                    segments = extract_segments(path)
                except NeedsOCR as e:
                    text = self._cached_ocr(rel, path)
                    if text is None and ocr is not None:
                        log(f"  👁 Leyendo con IA {rel}…")
                        text, cost = ocr(path)
                        stats.cost += cost
                        stats.ocr_done += 1
                        self.conn.execute("INSERT OR REPLACE INTO ocr VALUES (?, ?, ?)",
                                          (f"{rel}|{sig}", text, datetime.now().isoformat(timespec="seconds")))
                    if text is None:
                        segments, status, detail = [], "sin_texto", str(e)
                        stats.no_text.append(rel)
                    else:
                        segments, source = [("", text)], "ocr"
            except AIFatalError:
                raise  # clave no válida, sin conexión…: no tiene sentido seguir
            except Exception as e:  # un documento dañado no para el resto
                segments, status, detail = [], "error", f"{type(e).__name__}: {e}"
                stats.errors.append(f"{rel}: {detail}")
            self._delete(rel)
            pieces = chunk(segments)
            for loc, text in pieces:
                cur = self.conn.execute("INSERT INTO chunks (path, loc, text) VALUES (?, ?, ?)", (rel, loc, text))
                self.conn.execute("INSERT INTO chunks_fts (rowid, text) VALUES (?, ?)", (cur.lastrowid, text))
            self.conn.execute("INSERT OR REPLACE INTO docs VALUES (?, ?, ?, ?, ?, ?)",
                              (rel, sig, status, source, detail, datetime.now().isoformat(timespec="seconds")))
            self.conn.commit()
            if status == "ok":
                stats.indexed += 1
                log(f"  🔎 {rel} ({len(pieces)} fragmentos)")
        stats.semantic_error = self._embed_missing(log)
        self._matrix = None
        return stats

    def _delete(self, rel: str) -> None:
        rows = self.conn.execute("SELECT id, text FROM chunks WHERE path = ?", (rel,)).fetchall()
        for r in rows:
            self.conn.execute("INSERT INTO chunks_fts (chunks_fts, rowid, text) VALUES ('delete', ?, ?)",
                              (r["id"], r["text"]))
        self.conn.executemany("DELETE FROM vectors WHERE chunk_id = ?", [(r["id"],) for r in rows])
        self.conn.execute("DELETE FROM chunks WHERE path = ?", (rel,))
        self.conn.execute("DELETE FROM docs WHERE path = ?", (rel,))

    def _embed_missing(self, log) -> str | None:
        if not self.embedder.available():
            return self.embedder.error
        rows = self.conn.execute(
            "SELECT id, text FROM chunks WHERE id NOT IN (SELECT chunk_id FROM vectors)").fetchall()
        if not rows:
            return None
        log(f"  🧠 Preparando la búsqueda por significado ({len(rows)} fragmentos)…")
        try:
            for i in range(0, len(rows), 64):
                batch = rows[i:i + 64]
                vecs = self.embedder.embed([r["text"] for r in batch])
                self.conn.executemany("INSERT OR REPLACE INTO vectors VALUES (?, ?)",
                                      [(r["id"], v.tobytes()) for r, v in zip(batch, vecs)])
                self.conn.commit()
        except Exception as e:
            return self.embedder.error or str(e)
        return None

    # ---------------------------------------------------------- buscar
    def search(self, query: str, paths: set[str] | None = None, limit: int = 20) -> list[dict]:
        """Búsqueda híbrida: palabras + significado, combinadas por posición."""
        query = query.strip()
        if not query:
            return []
        keyword = self._keyword(query, paths, limit * 3)
        semantic = self._semantic(query, paths, limit * 3)
        scores: dict[int, float] = {}
        info: dict[int, dict] = {}
        for ranking, kind in ((keyword, "palabras"), (semantic, "significado")):
            for pos, r in enumerate(ranking):
                scores[r["id"]] = scores.get(r["id"], 0) + 1 / (60 + pos)  # fusión por rangos
                entry = info.setdefault(r["id"], {**r, "match": []})
                entry["match"].append(kind)
                if kind == "palabras":
                    entry["snippet"] = r["snippet"]
        ranked = sorted(scores, key=scores.get, reverse=True)
        out, seen = [], set()
        for cid in ranked:
            r = info[cid]
            key = (r["path"], r["loc"])
            if key in seen:  # un resultado por página/diapositiva
                continue
            seen.add(key)
            out.append({"path": r["path"], "loc": r["loc"], "snippet": r["snippet"],
                        "text": r["text"], "match": r["match"]})
            if len(out) >= limit:
                break
        return out

    @staticmethod
    def _fts_query(query: str, joiner: str) -> str:
        words = re.findall(r"\w+", query, flags=re.UNICODE)
        # Las palabras vacías («de», «la»…) no ayudan a encontrar nada.
        useful = [w for w in words if w.lower() not in STOPWORDS] or words
        return f" {joiner} ".join(f'"{w}"' for w in useful)

    def _keyword(self, query: str, paths, limit: int) -> list[dict]:
        rows = []
        for joiner in ("AND", "OR"):  # primero todas las palabras; si no hay nada, alguna
            q = self._fts_query(query, joiner)
            if not q:
                return []
            rows = self.conn.execute(
                f"""SELECT c.id, c.path, c.loc, c.text,
                           snippet(chunks_fts, 0, '{HL_OPEN}', '{HL_CLOSE}', '…', 24) AS snippet
                    FROM chunks_fts JOIN chunks c ON c.id = chunks_fts.rowid
                    WHERE chunks_fts MATCH ? ORDER BY bm25(chunks_fts) LIMIT ?""",
                (q, limit * 4)).fetchall()
            rows = [dict(r) for r in rows if paths is None or r["path"] in paths][:limit]
            if rows:
                break
        return rows

    def _semantic(self, query: str, paths, limit: int) -> list[dict]:
        if not self.embedder.available():
            return []
        import numpy as np

        if self._matrix is None:
            rows = self.conn.execute("SELECT chunk_id, vec FROM vectors").fetchall()
            if not rows:
                return []
            self._matrix = (np.array([r["chunk_id"] for r in rows]),
                            np.vstack([np.frombuffer(r["vec"], dtype="float32") for r in rows]))
        ids, matrix = self._matrix
        try:
            q = self.embedder.embed([query])[0]
        except Exception:
            return []
        sims = matrix @ q
        out = []
        for idx in np.argsort(-sims):
            if sims[idx] < MIN_SIMILARITY or len(out) >= limit:
                break
            r = self.conn.execute("SELECT id, path, loc, text FROM chunks WHERE id = ?",
                                  (int(ids[idx]),)).fetchone()
            if r and (paths is None or r["path"] in paths):
                text = r["text"]
                out.append({**dict(r), "snippet": text[:220] + ("…" if len(text) > 220 else "")})
        return out


OCR_PROMPT = """\
Transcribe todo el texto de este material de estudio (apuntes a mano, foto de
la pizarra, ejercicio o documento escaneado). Conserva la estructura (títulos,
listas, pasos) y escribe las fórmulas de forma legible, en notación LaTeX
sencilla si hace falta. No resumas ni corrijas: copia lo que pone. Si una parte
no se lee, escribe [ilegible]. Si no hay texto, responde exactamente: (sin texto)."""


def claude_ocr(client, model: str):
    """Devuelve una función ocr(path) -> (texto, coste) que usa Claude."""
    from .ai import AIError, create_message, material_block, usage_cost

    def ocr(path: Path) -> tuple[str, float]:
        resp = create_message(
            client, model, max_tokens=16000, thinking={"type": "adaptive"},
            output_config={"effort": "low"},
            messages=[{"role": "user", "content": [material_block(path, path.name),
                                                    {"type": "text", "text": OCR_PROMPT}]}],
        )
        if resp.stop_reason == "refusal":
            raise AIError("la IA no ha podido leer este archivo")
        text = "".join(b.text for b in resp.content if b.type == "text").strip()
        return ("" if text == "(sin texto)" else text), usage_cost(resp.model or model, resp.usage)
    return ocr
